"""Merge texts/raw/*.json into blinded text files plus a separate answer key.

texts/texts.json      - text_id -> text, per assignment, in shuffled order (no grade info)
texts/texts_key.json  - text_id -> intended grade, F type, rationale (keep away from graders)
texts/review_leonie.xlsx - blind review workbook for the expert check
"""
import json
import random
from pathlib import Path

ROOT = Path(__file__).parent
RAW = ROOT / "texts" / "raw"
SEED = 20260925
GRADES = ["F", "E", "D", "C", "B", "A"]

assignments = {a["id"]: a for a in json.load(open(ROOT / "assignments.json", encoding="utf-8"))["assignments"]}

texts, key = {}, {}
rng = random.Random(SEED)
for aid in assignments:
    raw = json.load(open(RAW / f"{aid}.json", encoding="utf-8"))["texts"]
    assert len(raw) == 60, (aid, len(raw))
    assert all(sum(t["intended_grade"] == g for t in raw) == 10 for g in GRADES), aid
    order = list(range(60))
    rng.shuffle(order)
    texts[aid] = []
    for pos, i in enumerate(order, start=1):
        t = raw[i]
        tid = f"{aid}-T{pos:02d}"
        texts[aid].append({"text_id": tid, "text": t["text"]})
        key[tid] = {
            "assignment_id": aid,
            "intended_grade": t["intended_grade"],
            "f_type": t.get("f_type"),
            "rationale": t.get("rationale"),
            "raw_n": t["n"],
        }

with open(ROOT / "texts" / "texts.json", "w", encoding="utf-8") as f:
    json.dump(texts, f, ensure_ascii=False, indent=2)
with open(ROOT / "texts" / "texts_key.json", "w", encoding="utf-8") as f:
    json.dump(key, f, ensure_ascii=False, indent=2)

# Blind review workbook
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
info = wb.active
info.title = "Instructions"
for row in [
    ["Blind review of synthetic learner texts (German, Swedish grundskola, grades 6-9)"],
    [],
    ["For each text, please enter:"],
    ["grade", "Your grade F/E/D/C/B/A using the task criteria (sheet 'Criteria')."],
    ["realistic", "1-5: how realistic is this as a text by a Swedish student in this grade? (5 = very realistic)"],
    ["neutrality_issue", "y if the text reveals anything about the writer's gender, origin, religion, home language or family structure; otherwise leave empty"],
    ["comment", "Anything else, e.g. unrealistic errors or a suggested fix."],
    [],
    ["The intended grades are kept separately so this review stays blind."],
]:
    info.append(row)
info["A1"].font = Font(bold=True, size=13)
info.column_dimensions["A"].width = 18
info.column_dimensions["B"].width = 110

crit = wb.create_sheet("Criteria")
crit.append(["assignment", "grade", "title", "instruction (sv)", "E", "C", "A"])
for a in assignments.values():
    crit.append([a["id"], a["grade"], a["title_de"], a["instruction_sv"],
                 a["criteria"]["E"], a["criteria"]["C"], a["criteria"]["A"]])
for col, w in zip("ABCDEFG", [11, 7, 28, 50, 60, 60, 60]):
    crit.column_dimensions[col].width = w
for row in crit.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
for c in crit[1]:
    c.font = Font(bold=True)

for aid, items in texts.items():
    a = assignments[aid]
    ws = wb.create_sheet(aid)
    ws.append([f"{aid} (åk {a['grade']}): {a['title_de']}"])
    ws.append([a["instruction_sv"]])
    ws.append([])
    ws.append(["text_id", "text", "grade", "realistic", "neutrality_issue", "comment"])
    ws["A1"].font = Font(bold=True, size=12)
    for c in ws[4]:
        c.font = Font(bold=True)
    for item in items:
        ws.append([item["text_id"], item["text"], None, None, None, None])
    last = ws.max_row
    dv_grade = DataValidation(type="list", formula1='"F,E,D,C,B,A"', allow_blank=True)
    dv_real = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    dv_neu = DataValidation(type="list", formula1='"y"', allow_blank=True)
    for dv, col in [(dv_grade, "C"), (dv_real, "D"), (dv_neu, "E")]:
        ws.add_data_validation(dv)
        dv.add(f"{col}5:{col}{last}")
    for col, w in zip("ABCDEF", [12, 90, 8, 10, 16, 50]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"

wb.save(ROOT / "texts" / "review_leonie.xlsx")
print("ok:", sum(len(v) for v in texts.values()), "texts")
